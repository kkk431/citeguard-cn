"""End-to-end tests for the FastAPI app.

The full sample paper is uploaded through ``/api/analyze``; the background job is
polled to completion and the resulting report structure is asserted. All outbound
HTTP is replaced by the deterministic Crossref double, so the suite is offline and
does not need ``DEEPSEEK_API_KEY`` (support degrades to ``not_checked``).
"""
from __future__ import annotations

import pytest
from fastapi.testclient import TestClient

from citeguard import metadata_verify as mv
from citeguard.reference_parser import parse_reference

from helpers import (
    ROOT,
    crossref_message,
    crossref_records,
    install_fake_get,
    make_doi_handler,
    wait_for_job,
)

SAMPLE = ROOT / "web" / "static" / "samples" / "sample_paper.txt"


@pytest.fixture
def client(isolated_project):
    import web.app as webapp

    return TestClient(webapp.app)


@pytest.fixture
def plain_client():
    import web.app as webapp

    return TestClient(webapp.app)


def test_health_endpoint(client):
    resp = client.get("/api/health")
    assert resp.status_code == 200
    payload = resp.json()
    assert payload["ok"] is True
    assert payload["version"]
    assert payload["offline"] is False
    assert {".pdf", ".docx", ".txt", ".md"} <= set(payload["supported"])


def test_index_page_renders(client):
    resp = client.get("/")
    assert resp.status_code == 200
    assert "引证卫士" in resp.text


def test_history_page_renders(client):
    resp = client.get("/history")
    assert resp.status_code == 200
    text = resp.text
    assert "任务历史" in text
    assert "sk-" not in text


def test_about_page_renders(client):
    resp = client.get("/about")
    assert resp.status_code == 200
    assert "CiteGuard" in resp.text


def test_benchmark_api_serves_shipped_results(plain_client):
    resp = plain_client.get("/api/benchmark")
    assert resp.status_code == 200
    payload = resp.json()
    assert payload["available"] is True
    assert payload["n"] == 250
    assert "Exact" in (payload.get("per_class") or {})


def test_benchmark_page_renders(plain_client):
    resp = plain_client.get("/benchmark")
    assert resp.status_code == 200
    assert "CiteBench" in resp.text
    assert "运行状态" in resp.text
    assert "标签质量审计" in resp.text
    assert "中文覆盖率试点" in resp.text
    assert "万方" in resp.text


def test_benchmark_status_never_leaks_secrets(plain_client, monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-secret-should-not-leak")
    resp = plain_client.get("/benchmark")
    assert resp.status_code == 200
    assert "运行状态" in resp.text
    assert "sk-secret-should-not-leak" not in resp.text


def test_analyze_requires_input(client):
    resp = client.post("/api/analyze")
    assert resp.status_code == 400


def test_analyze_rejects_unsupported_suffix(client):
    resp = client.post("/api/analyze", files={"file": ("evil.exe", b"MZ", "application/octet-stream")})
    assert resp.status_code == 400


def test_analyze_rejects_oversized_upload(client, monkeypatch, isolated_project):
    import web.app as webapp

    monkeypatch.setattr(webapp, "MAX_UPLOAD_BYTES", 8)
    resp = client.post(
        "/api/analyze",
        files={"file": ("big.txt", b"0123456789" * 4, "text/plain")},
    )
    assert resp.status_code == 413
    assert not list((isolated_project / "uploads").iterdir())


def test_unknown_job_and_report_return_404(client):
    assert client.get("/api/jobs/nope").status_code == 404
    assert client.get("/api/report/nope").status_code == 404
    assert client.get("/report/nope").status_code == 404
    assert client.get("/api/report/nope/export").status_code == 404


def test_e2e_upload_sample_paper(client, monkeypatch, isolated_project, sample_refs):
    records = crossref_records(sample_refs)
    calls = install_fake_get(monkeypatch, mv, make_doi_handler(records))

    with SAMPLE.open("rb") as fh:
        resp = client.post(
            "/api/analyze",
            files={"file": (SAMPLE.name, fh, "text/plain")},
        )
    assert resp.status_code == 202
    job_id = resp.json()["job_id"]

    job = wait_for_job(client, job_id)
    assert job["status"] == "done", job
    assert job["progress"] == 100
    assert job["filename"] == SAMPLE.name
    assert job["error"] == ""

    resp = client.get(f"/api/report/{job_id}")
    assert resp.status_code == 200
    report = resp.json()

    assert report["job_id"] == job_id
    stats = report["stats"]
    assert stats["total"] == 12
    assert stats["intext_citations"] == 12
    assert stats["Exact"] == 8
    assert stats["Minor"] == 2
    assert stats["Major"] == 2
    assert stats["Unknown"] == 0
    assert stats["Exact"] + stats["Minor"] + stats["Major"] + stats["Unknown"] == stats["total"]
    assert 0 <= stats["avg_compliance"] <= 100
    assert 0 <= stats["risk_index"] <= 100
    # no key -> the judge is skipped and everything is reported as not_checked
    assert stats["support"]["not_checked"] == 12

    assert len(report["references"]) == 12
    expected_keys = {"index", "raw", "parsed", "verify", "compliance", "support"}
    for item in report["references"]:
        assert expected_keys <= set(item)
        assert item["verify"]["status"] in {"Exact", "Minor", "Major", "Unknown"}
        assert "score" in item["compliance"]
    assert calls, "the fake HTTP transport should have been exercised"
    assert {"judge_available", "cap", "candidates"} <= set(stats["support_meta"])
    assert stats["support_meta"]["judge_available"] is False

    # uploads are deleted after extraction (privacy)
    assert not list((isolated_project / "uploads").iterdir())

    # report HTML page
    page = client.get(f"/report/{job_id}")
    assert page.status_code == 200
    assert job_id in page.text
    assert "引证卫士" in page.text

    # exports
    md = client.get(f"/api/report/{job_id}/export?format=md")
    assert md.status_code == 200
    assert "markdown" in md.headers["content-type"]
    assert "# 引证卫士" in md.text

    as_json = client.get(f"/api/report/{job_id}/export?format=json")
    assert as_json.status_code == 200
    assert as_json.json()["job_id"] == job_id

    html = client.get(f"/api/report/{job_id}/export?format=html")
    assert html.status_code == 200
    assert "<!doctype html>" in html.text

    refs = client.get(f"/api/report/{job_id}/export?format=refs")
    assert refs.status_code == 200
    assert "text/plain" in refs.headers["content-type"]
    assert refs.text.startswith("[1] ")

    ris = client.get(f"/api/report/{job_id}/export?format=ris")
    assert ris.status_code == 200
    assert "x-research-info-systems" in ris.headers["content-type"]
    assert "TY  - " in ris.text and "ER  - " in ris.text

    enw = client.get(f"/api/report/{job_id}/export?format=enw")
    assert enw.status_code == 200
    assert "%0 " in enw.text

    import io
    import zipfile

    bundle = client.get(f"/api/report/{job_id}/bundle")
    assert bundle.status_code == 200
    with zipfile.ZipFile(io.BytesIO(bundle.content)) as zf:
        names = set(zf.namelist())
    assert {"report.json", "report.html", "report.ris", "report.enw", "manifest.json"} <= names

    assert client.get(f"/api/report/{job_id}/export?format=pdf").status_code == 400


def test_e2e_pasted_reference_text(client, monkeypatch):
    ref = parse_reference(
        "[1] 张三. 题名[J]. 期刊, 2020, 31(2): 10-20. DOI: 10.1234/paste."
    )
    install_fake_get(monkeypatch, mv, make_doi_handler({ref.doi: crossref_message(ref)}))
    pasted = (
        "[1] 张三. 题名[J]. 期刊, 2020, 31(2): 10-20. DOI: 10.1234/paste.\n"
    )

    resp = client.post("/api/analyze", data={"text": pasted})
    assert resp.status_code == 202
    job_id = resp.json()["job_id"]
    job = wait_for_job(client, job_id)
    assert job["status"] == "done", job

    report = client.get(f"/api/report/{job_id}").json()
    assert report["stats"]["total"] == 1
    assert report["stats"]["Exact"] == 1
    assert report["input"] == "粘贴文本"


def test_e2e_upload_docx(client, monkeypatch, tmp_path):
    from docx import Document

    refs = [
        parse_reference("[1] 张三. 题名一[J]. 期刊, 2020, 31(2): 10-20. DOI: 10.1234/d1."),
        parse_reference("[2] 李四. 题名二[J]. 期刊, 2021, 65(3): 5-9. DOI: 10.1234/d2."),
    ]
    doc = Document()
    doc.add_paragraph("正文引用[1]与[2]。")
    doc.add_paragraph("参考文献")
    for ref in refs:
        doc.add_paragraph(ref.raw)
    path = tmp_path / "paper.docx"
    doc.save(path)

    records = {ref.doi: crossref_message(ref) for ref in refs}
    install_fake_get(monkeypatch, mv, make_doi_handler(records))

    with path.open("rb") as fh:
        resp = client.post(
            "/api/analyze",
            files={"file": (path.name, fh, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert resp.status_code == 202
    job_id = resp.json()["job_id"]
    job = wait_for_job(client, job_id)
    assert job["status"] == "done", job

    report = client.get(f"/api/report/{job_id}").json()
    assert report["stats"]["total"] == 2
    assert report["stats"]["Exact"] == 2
    assert report["stats"]["intext_citations"] == 2


def test_rejected_upload_does_not_create_job(client, isolated_project):
    import citeguard.jobs as jobs

    before = len(jobs.manager._jobs)
    client.post("/api/analyze", files={"file": ("evil.exe", b"MZ", "application/octet-stream")})
    assert len(jobs.manager._jobs) == before

def test_compare_page_renders(plain_client):
    resp = plain_client.get("/compare")
    assert resp.status_code == 200
    assert "多篇文献对比" in resp.text


def test_compare_api_rejects_without_judge(plain_client, monkeypatch):
    from citeguard import compare

    monkeypatch.setattr(compare, "available", lambda: False)
    resp = plain_client.post(
        "/api/compare", files=[("files", ("a.txt", b"x", "text/plain"))]
    )
    assert resp.status_code == 503


def test_compare_api_validates_file_count(client, monkeypatch):
    from citeguard import compare

    monkeypatch.setattr(compare, "available", lambda: True)
    resp = client.post(
        "/api/compare", files=[("files", ("a.txt", b"x", "text/plain"))]
    )
    assert resp.status_code == 400


def test_compare_job_end_to_end(client, monkeypatch):
    from citeguard import compare

    monkeypatch.setattr(compare, "available", lambda: True)

    def fake_extract(text, *, use_cache=True):
        return {
            "fields": {key: f"{key}-内容" for key in compare.FIELD_KEYS},
            "model": "fake",
            "excerpt_chars": len(text),
            "latency_ms": 1,
            "cached": False,
        }

    monkeypatch.setattr(compare, "extract_profile", fake_extract)
    files = [
        ("files", ("paper_a.txt", "标题A\n摘要：研究问题A。".encode("utf-8"), "text/plain")),
        ("files", ("paper_b.txt", "标题B\n摘要：研究问题B。".encode("utf-8"), "text/plain")),
    ]
    resp = client.post("/api/compare", files=files)
    assert resp.status_code == 202, resp.text
    job_id = resp.json()["job_id"]
    job = wait_for_job(client, job_id)
    assert job["status"] == "done", job

    data = client.get(f"/api/compare/{job_id}").json()
    result = data["result"]
    assert [paper["name"] for paper in result["papers"]] == ["paper_a", "paper_b"]
    assert result["papers"][0]["fields"]["method"] == "method-内容"

    md = client.get(f"/api/compare/{job_id}/export?format=md")
    assert md.status_code == 200
    assert "多篇文献对比" in md.text
    csv_resp = client.get(f"/api/compare/{job_id}/export?format=csv")
    assert csv_resp.status_code == 200
    assert "字段" in csv_resp.text
    bad = client.get(f"/api/compare/{job_id}/export?format=bib")
    assert bad.status_code == 400


def test_discover_page_renders(plain_client):
    resp = plain_client.get("/discover")
    assert resp.status_code == 200
    assert "discover-form" in resp.text
    assert "discover-end" in resp.text
    assert "全选已加载" in resp.text


def test_discover_api_rejects_short_query(plain_client):
    resp = plain_client.post("/api/discover", json={"query": "a"})
    assert resp.status_code in (400, 422)


def test_discover_api_returns_results(plain_client, monkeypatch):
    import web.app as webapp
    def fake_search(query, **kwargs):
        assert query == "ai education"
        assert kwargs["limit"] == 5
        return {
            "query": query,
            "matched": 2,
            "results": [
                {"title": "Paper A", "authors": ["A"], "year": 2020, "venue": "V",
                 "doi": "10.1/a", "cited_by": 3, "is_oa": False, "oa_url": "",
                 "openalex_id": "W1"},
                {"title": "Paper B", "authors": [], "year": 2021, "venue": "",
                 "doi": "", "cited_by": 0, "is_oa": True, "oa_url": "https://x",
                 "openalex_id": "W2"},
            ],
        }

    monkeypatch.setattr(webapp.discover, "search_papers", fake_search)
    resp = plain_client.post("/api/discover", json={"query": "ai education", "limit": 5})
    assert resp.status_code == 200, resp.text
    payload = resp.json()
    assert payload["matched"] == 2
    assert [item["title"] for item in payload["results"]] == ["Paper A", "Paper B"]


def test_discover_api_maps_upstream_error(plain_client, monkeypatch):
    import web.app as webapp
    def boom(query, **kwargs):
        raise webapp.discover.DiscoverError("OpenAlex 请求失败：500")

    monkeypatch.setattr(webapp.discover, "search_papers", boom)
    resp = plain_client.post("/api/discover", json={"query": "test topic"})
    assert resp.status_code == 502
    assert "OpenAlex" in resp.json()["detail"]


def test_discover_import_returns_counts(plain_client, monkeypatch):
    import web.app as webapp
    monkeypatch.setattr(
        webapp.discover,
        "import_entries",
        lambda items: {"added": len(items), "skipped": 1, "count": 7},
    )
    resp = plain_client.post(
        "/api/discover/import",
        json={"items": [{"title": "Paper A", "year": 2020}]},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["added"] == 1
    assert resp.json()["count"] == 7


def test_discover_import_rejects_empty(plain_client):
    resp = plain_client.post("/api/discover/import", json={"items": []})
    assert resp.status_code == 400


def test_discover_export_csv(plain_client, monkeypatch):
    import web.app as webapp
    monkeypatch.setattr(
        webapp.discover,
        "search_papers",
        lambda query, **kwargs: {"query": query, "matched": 1, "results": []},
    )
    monkeypatch.setattr(webapp.discover, "render_csv", lambda result: "title,doi\nA,\n")
    resp = plain_client.get("/api/discover/export?query=test")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")


def test_access_token_guard(plain_client, monkeypatch):
    import web.app as webapp

    monkeypatch.setattr(webapp.config, "ACCESS_TOKEN", "s3cret")
    plain_client.cookies.clear()
    assert plain_client.get("/api/library").status_code == 401
    assert (
        plain_client.get("/api/library", headers={"X-Access-Token": "s3cret"}).status_code == 200
    )
    # Opening a page with ?token= stores an HttpOnly cookie for later fetches.
    plain_client.cookies.clear()
    assert plain_client.get("/?token=s3cret").status_code == 200
    assert plain_client.get("/api/library").status_code == 200
    plain_client.cookies.clear()
    assert plain_client.get("/api/library").status_code == 401


def test_rate_limit_blocks_post_floods(plain_client, monkeypatch):
    import web.app as webapp

    monkeypatch.setattr(webapp.config, "ACCESS_TOKEN", "")
    monkeypatch.setattr(webapp.config, "RATE_LIMIT_PER_MIN", 2)
    webapp._rate_buckets.clear()
    # Empty library -> handler answers 400, but every POST still consumes a slot.
    for _ in range(2):
        assert plain_client.post("/api/library/enrich").status_code == 400
    resp = plain_client.post("/api/library/enrich")
    assert resp.status_code == 429
    assert resp.headers.get("Retry-After") == "60"

def test_library_add_accepts_manual_entry(client):
    payload = {
        "title": "中文期刊题录手工补录测试",
        "authors": "张三；李四",
        "year": "2024",
        "venue": "某中文期刊",
    }
    resp = client.post("/api/library/add", json=payload)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["added"] == 1 and body["count"] == 1

    again = client.post("/api/library/add", json=payload)
    assert again.status_code == 200
    assert again.json()["added"] == 0  # same title deduped

    listing = client.get("/api/library").json()
    assert listing["count"] == 1
    assert listing["entries"][0]["title"] == payload["title"]


def test_library_bulk_accepts_multiple_lines(client):
    text = "\n".join(
        [
            "[1] 张三. 中文引文核验方法[J]. 情报学报, 2024, 12(3): 10-20.",
            "[2] 李四. 另一篇中文论文[J]. 图书馆论坛, 2023, 43(2): 1-9.",
            "...",
        ]
    )
    resp = client.post("/api/library/bulk", json={"text": text})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["added"] == 2
    assert body["count"] == 2
    listing = client.get("/api/library").json()
    assert listing["count"] == 2
    titles = " ".join(e["title"] for e in listing["entries"])
    assert "中文引文核验方法" in titles and "另一篇中文论文" in titles

def test_compare_reports_positional_deltas():
    import web.app as webapp

    a = {
        "input": "old.pdf",
        "generated_at": "2026-10-27 10:00",
        "stats": {"total": 2, "Exact": 1, "Minor": 0, "Major": 1, "Unverified": 0,
                  "avg_compliance": 80, "risk_index": 30},
        "references": [
            {"parsed": {"title": "A"}, "verify": {"status": "Exact"}},
            {"parsed": {"title": "B"}, "verify": {"status": "Major"}},
        ],
    }
    b = {
        "input": "new.pdf",
        "generated_at": "2026-10-27 11:00",
        "stats": {"total": 2, "Exact": 2, "Minor": 0, "Major": 0, "Unverified": 0,
                  "avg_compliance": 95, "risk_index": 10},
        "references": [
            {"parsed": {"title": "A"}, "verify": {"status": "Exact"}},
            {"parsed": {"title": "B"}, "verify": {"status": "Exact"}},
        ],
    }
    out = webapp._compare_reports(a, b)
    assert out["stats"]["Major"]["delta"] == -1
    assert out["stats"]["avg_compliance"]["delta"] == 15
    assert out["changed"] == 1
    assert out["changes"][0]["from"] == "Major"
    assert out["changes"][0]["to"] == "Exact"
    assert out["positional"] is True


def test_compare_api_requires_two_reports(plain_client):
    r = plain_client.get("/api/compare?a=nope&b=missing")
    assert r.status_code == 404

def test_rerun_submits_retained_text(plain_client, monkeypatch):
    import web.app as webapp

    job = webapp.manager.create(filename="old.pdf")
    d = webapp.config.DATA_DIR / "jobs" / job["id"]
    d.mkdir(parents=True, exist_ok=True)
    (d / "input.txt").write_text("保留的正文与参考文献", encoding="utf-8")
    (d / "options.json").write_text('{"support_max": 7}', encoding="utf-8")

    captured = {}
    monkeypatch.setattr(webapp.manager, "submit", lambda jid, fn, *a, **k: captured.update(jid=jid, args=a, kwargs=k))
    r = plain_client.post(f"/api/jobs/{job['id']}/rerun")
    assert r.status_code == 202
    body = r.json()
    assert body["rerun_of"] == job["id"]
    assert captured["jid"] == body["job_id"]
    assert captured["kwargs"]["text"] == "保留的正文与参考文献"
    assert captured["kwargs"]["support_max"] == 7


def test_rerun_without_retained_text_returns_409(plain_client):
    import web.app as webapp

    job = webapp.manager.create(filename="gone.pdf")
    r = plain_client.post(f"/api/jobs/{job['id']}/rerun")
    assert r.status_code == 409


def test_stale_running_job_becomes_interrupted(plain_client):
    import web.app as webapp

    job = webapp.manager.create(filename="crashed.pdf")
    webapp.manager.update(job["id"], status="running", progress=30)
    changed = webapp.manager.mark_stale_as_interrupted()
    assert changed >= 1
    again = webapp.manager.get(job["id"])
    assert again["status"] == "interrupted"
    assert "重启" in again["error"]
