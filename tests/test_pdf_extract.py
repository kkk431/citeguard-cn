"""Tests for text/PDF extraction and body/reference splitting."""
from __future__ import annotations

import pytest

from citeguard.pdf_extract import (
    extract_page_text,
    extract_text,
    extract_text_with_meta,
    figure_number_from_claim,
    find_figure_page,
    normalize_text,
    split_references,
)


def test_normalize_text_collapses_whitespace():
    assert normalize_text("a\u00a0 b\t\tc\u3000d") == "a b c d"


def test_normalize_text_caps_blank_lines():
    assert normalize_text("a\n\n\n\nb") == "a\n\nb"


def test_split_references_chinese_heading():
    text = "正文第一段[1]。\n\n## 参考文献\n[1] 张三. 题名[J]. 期刊, 2020, 1(2): 3-4.\n"
    body, refs = split_references(text)
    assert "正文第一段" in body
    assert "参考文献" not in refs
    assert refs.strip().startswith("[1] 张三")


def test_split_references_english_heading():
    body, refs = split_references("Body text.\n\nReferences\n[1] A. B. Title[J]. J, 2020.\n")
    assert body.strip() == "Body text."
    assert "[1] A. B. Title" in refs


def test_split_references_markdown_bullet_heading():
    body, refs = split_references("# Title\n\n- References\n[1] A. T[J]. J, 2020.\n")
    assert "# Title" in body
    assert refs.strip().startswith("[1]")


def test_split_references_without_heading_returns_whole_text():
    text = "no heading here, just text"
    assert split_references(text) == (text, "")


def test_extract_text_reads_txt(tmp_path):
    p = tmp_path / "paper.txt"
    p.write_text("Hello\u3000World", encoding="utf-8")
    assert extract_text(p) == "Hello World"


def test_extract_text_reads_markdown(tmp_path):
    p = tmp_path / "paper.md"
    p.write_text("# 标题\n\n正文", encoding="utf-8")
    assert "# 标题" in extract_text(p)


def test_extract_text_rejects_unknown_suffix(tmp_path):
    p = tmp_path / "paper.rtf"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        extract_text(p)


def test_extract_text_rejects_legacy_doc_with_hint(tmp_path):
    p = tmp_path / "old.doc"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError) as excinfo:
        extract_text(p)
    assert ".docx" in str(excinfo.value)


def test_extract_text_reads_docx_paragraphs_and_tables(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一段正文")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "参考文献"
    table.rows[0].cells[1].text = "[1] 张三. 题名[J]. 期刊, 2020."
    doc.add_paragraph("尾段正文")
    path = tmp_path / "paper.docx"
    doc.save(path)

    text = extract_text(path)

    assert "第一段正文" in text
    assert "尾段正文" in text
    assert "参考文献" in text
    assert "[1] 张三" in text


def test_extract_text_unreadable_pdf_raises_runtime_error(tmp_path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"%PDF-1.4\nthis is not a real pdf body")
    with pytest.raises(RuntimeError):
        extract_text(p)


def test_utf8_bom_does_not_drop_first_reference(tmp_path):
    """Regression: a BOM-prefixed .txt must keep its first [1] entry."""
    from citeguard.reference_parser import parse_reference_list

    body = "[1] 张三. 题名一[J]. 期刊, 2020.\n[2] 李四. 题名二[J]. 期刊, 2021."
    path = tmp_path / "bom.txt"
    path.write_bytes(b"\xef\xbb\xbf" + body.encode("utf-8"))

    refs = parse_reference_list(extract_text(path))

    assert len(refs) == 2
    assert refs[0].index == 1


def test_ocr_standalone_reference_numbers_are_merged():
    """OCR often puts "[n]" on its own line; merge it into the entry."""
    from citeguard.ocr import _merge_reference_numbers

    text = (
        "正文引用[1]。\n\n参考文献\n"
        "[1]\n张三. 题名[J]. 期刊, 2020, 1(2): 3-4.\n"
        "[2]\n李四. 题名二[J]. 期刊, 2021, 2(1): 5-6.\n"
    )

    merged = _merge_reference_numbers(text)

    assert "[1] 张三. 题名[J]." in merged
    assert "[2] 李四. 题名二[J]." in merged
    assert "正文引用[1]。" in merged  # body untouched


def test_looks_garbled_detects_cid_fonts():
    from citeguard.pdf_extract import _looks_garbled

    assert _looks_garbled("(cid:131)(cid:243)" * 20)
    assert not _looks_garbled("正常的中文文本，参考文献[J]。")
    assert not _looks_garbled("")
    # Some embedded CJK fonts decode into raw C0 control bytes, not (cid:...) tokens.
    assert _looks_garbled("ADMM" + "\x15\x1b\x11\x12" * 40 + "text")
    assert not _looks_garbled("normal line one\nline two\twith tab")


def test_garbled_pdf_falls_back_to_ocr(monkeypatch, tmp_path):
    import citeguard.pdf_extract as px

    pdf = tmp_path / "garbled.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")
    monkeypatch.setattr(px, "_pdf_pdfplumber", lambda path: "(cid:131)(cid:243)" * 50)
    monkeypatch.setattr(px, "_pdf_pypdf", lambda path: "")
    monkeypatch.setattr(px, "_ocr_pdf", lambda path: "参考文献\n[1] 张三. 题名[J]. 期刊, 2020.")

    text = px.extract_text(pdf)

    assert "张三" in text and "cid:" not in text


def test_extract_text_with_meta_reports_plain_text(tmp_path):
    p = tmp_path / "paper.txt"
    p.write_text("Hello World", encoding="utf-8")

    text, meta = extract_text_with_meta(p)

    assert text == "Hello World"
    assert meta["method"] == "text"
    assert meta["ocr"] is False
    assert meta["chars"] == len("Hello World")
    assert meta["warnings"] == []


def test_extract_text_with_meta_marks_garbled_ocr_fallback(monkeypatch, tmp_path):
    import citeguard.pdf_extract as px

    pdf = tmp_path / "garbled.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")
    monkeypatch.setattr(px, "_pdf_page_count", lambda path: 3)
    monkeypatch.setattr(px, "_pdf_pdfplumber", lambda path: "(cid:131)(cid:243)" * 50)
    monkeypatch.setattr(px, "_pdf_pypdf", lambda path: "")
    monkeypatch.setattr(px, "_ocr_pdf", lambda path: "参考文献\n[1] 张三. 题名[J]. 期刊, 2020.")

    text, meta = px.extract_text_with_meta(pdf)

    assert "张三" in text
    assert meta["method"] == "ocr"
    assert meta["ocr"] is True
    assert meta["pages"] == 3
    assert any("文字层损坏" in w for w in meta["warnings"])
    assert not any("没有可用的文字层" in w for w in meta["warnings"])


def test_extract_text_with_meta_marks_scan_without_text_layer(monkeypatch, tmp_path):
    import citeguard.pdf_extract as px

    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")
    monkeypatch.setattr(px, "_pdf_page_count", lambda path: 2)
    monkeypatch.setattr(px, "_pdf_pdfplumber", lambda path: "")
    monkeypatch.setattr(px, "_pdf_pypdf", lambda path: "")
    monkeypatch.setattr(px, "_ocr_pdf", lambda path: "参考文献\n[1] 张三. 题名[J]. 期刊, 2020.")

    text, meta = px.extract_text_with_meta(pdf)

    assert "张三" in text
    assert meta["method"] == "ocr"
    assert any("没有可用的文字层" in w for w in meta["warnings"])


def test_report_exports_state_ocr_input():
    from citeguard.report import render_html, render_markdown

    report = {
        "generated_at": "2026-10-24 12:00:00",
        "input": "scan.pdf",
        "stats": {
            "total": 1,
            "Exact": 0,
            "Minor": 0,
            "Major": 0,
            "Unverified": 1,
            "Unknown": 0,
            "intext_citations": 0,
            "avg_compliance": 80,
            "risk_index": 30,
            "references": [],
        },
        "references": [],
        "formatted_reference_list": [],
        "input_meta": {
            "method": "ocr",
            "file_type": "pdf",
            "pages": 22,
            "ocr": True,
            "chars": 31000,
            "warnings": ["本文件没有可用的文字层，已用 OCR 识别 22 页"],
        },
        "disclaimer": "仅供参考。",
    }

    md = render_markdown(report)
    page = render_html(report)

    assert "文字来源：OCR 识别（扫描件） · 22 页 · 31000 字符" in md
    assert "文字来源提示" in page
    assert "OCR" in page


def test_garbled_pdf_without_ocr_raises_clear_error(monkeypatch, tmp_path):
    import citeguard.pdf_extract as px

    pdf = tmp_path / "garbled.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")
    monkeypatch.setattr(px, "_pdf_pdfplumber", lambda path: "(cid:131)(cid:243)" * 50)
    monkeypatch.setattr(px, "_pdf_pypdf", lambda path: "")
    monkeypatch.setattr(px, "_ocr_pdf", lambda path: "")

    with pytest.raises(RuntimeError) as excinfo:
        px.extract_text(pdf)
    assert "CID" in str(excinfo.value)


def test_ocr_failure_is_reported_instead_of_no_text_layer(monkeypatch, tmp_path):
    """A broken OCR environment must surface the real error, not "scanned PDF"."""
    import citeguard.pdf_extract as px

    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")
    monkeypatch.setattr(px, "_pdf_page_count", lambda path: 7)
    monkeypatch.setattr(px, "_pdf_pdfplumber", lambda path: "")
    monkeypatch.setattr(px, "_pdf_pypdf", lambda path: "")

    def broken_ocr(path):
        raise RuntimeError("OpenCV bindings requires numpy package")

    monkeypatch.setattr(px, "_ocr_pdf", broken_ocr)

    with pytest.raises(RuntimeError) as excinfo:
        px.extract_text(pdf)
    message = str(excinfo.value)
    assert "OCR 回退也失败" in message
    assert "OpenCV bindings requires numpy package" in message


def test_missing_ocr_dependency_is_explained(monkeypatch, tmp_path):
    import citeguard.pdf_extract as px

    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake")
    monkeypatch.setattr(px, "_pdf_page_count", lambda path: 2)
    monkeypatch.setattr(px, "_pdf_pdfplumber", lambda path: "")
    monkeypatch.setattr(px, "_pdf_pypdf", lambda path: "")

    def missing_dependency(path):
        raise RuntimeError(
            "扫描件 OCR 需要可选依赖 rapidocr-onnxruntime，请先执行 pip install rapidocr-onnxruntime"
        )

    monkeypatch.setattr(px, "_ocr_pdf", missing_dependency)

    with pytest.raises(RuntimeError) as excinfo:
        px.extract_text(pdf)
    assert "未安装 OCR 依赖" in str(excinfo.value)

def _blank_pdf(tmp_path, pages=1):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    path = tmp_path / "blank.pdf"
    with path.open("wb") as fh:
        writer.write(fh)
    return path


def test_extract_page_text_rejects_non_pdf(tmp_path):
    p = tmp_path / "x.txt"
    p.write_text("hi", encoding="utf-8")
    with pytest.raises(ValueError):
        extract_page_text(p, 1)


def test_extract_page_text_reports_blank_page(tmp_path):
    pdf = _blank_pdf(tmp_path, pages=2)
    assert pdf.exists()
    with pytest.raises(ValueError) as excinfo:
        extract_page_text(pdf, 1)
    assert "没有可提取的文字" in str(excinfo.value)


def test_extract_page_text_rejects_page_out_of_range(tmp_path):
    pdf = _blank_pdf(tmp_path, pages=1)
    with pytest.raises(ValueError) as excinfo:
        extract_page_text(pdf, 2)
    assert "超出范围" in str(excinfo.value)


def test_figure_number_from_claim():
    assert figure_number_from_claim("如图 3 所示，方法1 最好") == "3"
    assert figure_number_from_claim("Figure 4a shows the trend") == "4a"
    assert figure_number_from_claim("表 2 给出了样本量") == ""
    assert figure_number_from_claim("") == ""


def test_find_figure_page_matches_claim_number(tmp_path, monkeypatch):
    import sys
    import types

    from citeguard import pdf_extract as px

    class _Page:
        def __init__(self, text):
            self.text = text

    class _Doc:
        def __init__(self, pages):
            self.pages = pages

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

    fake_module = types.ModuleType("pdfplumber")
    fake_module.open = lambda path: _Doc(
        [_Page("引言段落，没有图注。"), _Page("图1 方法1 的准确率（%）\n正文继续")]
    )
    monkeypatch.setitem(sys.modules, "pdfplumber", fake_module)
    monkeypatch.setattr(px, "_page_text", lambda page: page.text)
    pdf = tmp_path / "fake.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    found = find_figure_page(pdf, "1")
    assert found is not None
    assert found[0] == 2
    assert found[1].startswith("图1")
    assert find_figure_page(pdf, "9") is None


def test_find_figure_page_scanned_pdf_returns_none(tmp_path):
    pdf = _blank_pdf(tmp_path, pages=2)
    assert find_figure_page(pdf, "1") is None
