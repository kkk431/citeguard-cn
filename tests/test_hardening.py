"""Hardening behaviours added during product work: input formats, list splitting,
job hygiene, support cap and upload validation.
"""
from __future__ import annotations

import io
import json
import time

import docx
import pytest
from fastapi.testclient import TestClient

import citeguard.jobs as jobs
import citeguard.metadata_verify as mv
from citeguard.pdf_extract import split_references
from citeguard.reference_parser import parse_reference_list, split_entries
from helpers import (
    crossref_records,
    install_fake_get,
    make_doi_handler,
    wait_for_job,
)

SAMPLE = (
    "本文讨论引文核验[1]。\n"
    "参考文献\n"
    "[1] He K, Zhang X, Ren S, et al. Deep residual learning for image "
    "recognition[C]. Proceedings of the IEEE Conference on Computer Vision "
    "and Pattern Recognition, 2016: 770-778. https://doi.org/10.1109/CVPR.2016.90\n"
)


# ---------------------------------------------------------------------------
# Reference-list splitting
# ---------------------------------------------------------------------------

def test_inline_reference_list_without_line_breaks_is_split():
    text = (
        "[1] He K, Zhang X. Deep residual learning for image recognition[C]. "
        "CVPR, 2016: 770-778. [2] Vaswani A, Shazeer N. Attention is all you need[C]. "
        "NeurIPS, 2017: 5998-6008. [3] 张三. 题名[J]. 期刊, 2020, 12(3): 45-56."
    )

    entries = split_entries(text)
    refs = parse_reference_list(text)

    assert len(entries) == 3
    assert [r.index for r in refs] == [1, 2, 3]
    assert refs[2].title == "题名"


def test_inline_split_keeps_brackets_inside_a_title():
    text = (
        "[1] 张三. 基于[图像]识别的系统[J]. 计算机学报, 2020, 12(3): 45-56. "
        "[2] 李四. 另一篇[J]. 软件学报, 2021, 3(1): 1-9."
    )

    entries = split_entries(text)

    assert len(entries) == 2
    assert "[图像]" in entries[0]


# ---------------------------------------------------------------------------
# Job hygiene
# ---------------------------------------------------------------------------

def test_cleanup_removes_old_job_records_and_report_dirs(tmp_path, monkeypatch):
    jobs_dir = tmp_path / "jobs"
    jobs_dir.mkdir()
    monkeypatch.setattr(jobs, "JOBS_DIR", jobs_dir)
    monkeypatch.setattr(jobs.manager, "_jobs", {})

    fresh = jobs_dir / "fresh.json"
    old = jobs_dir / "old.json"
    fresh.write_text(json.dumps({"id": "fresh"}), encoding="utf-8")
    old.write_text(json.dumps({"id": "old"}), encoding="utf-8")
    old_report = jobs_dir / "old"
    old_report.mkdir()
    (old_report / "report.json").write_text("{}", encoding="utf-8")

    past = time.time() - 30 * 86400
    import os

    os.utime(old, (past, past))

    removed = jobs.manager.cleanup(keep_recent=10, max_age_days=7)

    assert removed["records"] >= 1
    assert not old.exists()
    assert fresh.exists()
    assert not old_report.exists()


# ---------------------------------------------------------------------------
# Support cap
# ---------------------------------------------------------------------------

def test_support_cap_zero_marks_all_not_checked(isolated_project, monkeypatch, tmp_path):
    from citeguard.pipeline import run_job

    refs = parse_reference_list(SAMPLE.split("参考文献", 1)[1])
    records = crossref_records(refs)
    install_fake_get(monkeypatch, mv, make_doi_handler(records))

    job = jobs.manager.create(filename="cap0.txt")
    report = run_job(job["id"], text=SAMPLE, support_max=0)

    meta = report["stats"]["support_meta"]
    assert meta["cap"] == 0
    assert meta["judged"] == 0
    assert report["stats"]["support"]["not_checked"] == report["stats"]["total"]


def test_support_meta_reports_candidate_counts(isolated_project, monkeypatch):
    from citeguard.pipeline import run_job

    refs = parse_reference_list(SAMPLE.split("参考文献", 1)[1])
    install_fake_get(monkeypatch, mv, make_doi_handler(crossref_records(refs)))

    job = jobs.manager.create(filename="cap.txt")
    report = run_job(job["id"], text=SAMPLE, support_max=30)

    meta = report["stats"]["support_meta"]
    assert meta["cap"] == 30
    assert meta["candidates"] >= 1
    assert meta["judge_available"] is False  # conftest clears the API key


# ---------------------------------------------------------------------------
# Upload handling
# ---------------------------------------------------------------------------

@pytest.fixture
def client(isolated_project, monkeypatch):
    import web.app as webapp

    install_fake_get(monkeypatch, mv, _not_found)
    return TestClient(webapp.app)


def _not_found(url: str, params: dict):
    from helpers import FakeResponse

    return FakeResponse(404, {"status": "not found"})


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("测试正文[1]。")
    document.add_paragraph("参考文献")
    document.add_paragraph(
        "[1] He K, Zhang X. Deep residual learning for image recognition[C]. "
        "CVPR, 2016: 770-778."
    )
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_docx_upload_is_accepted_and_analyzed(client):
    resp = client.post(
        "/api/analyze",
        files={
            "file": (
                "paper.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert resp.status_code == 202, resp.text
    job = wait_for_job(client, resp.json()["job_id"])
    assert job["status"] == "done", job
    report = client.get(f"/api/report/{job['id']}").json()
    assert report["stats"]["total"] == 1


def test_empty_upload_is_rejected_without_leaving_a_job(client):
    before = len(jobs.manager._jobs)

    resp = client.post("/api/analyze", files={"file": ("empty.pdf", b"", "application/pdf")})

    assert resp.status_code == 400
    assert len(jobs.manager._jobs) == before


def test_health_reports_supported_suffixes(client):
    payload = client.get("/api/health").json()

    assert payload["ok"] is True
    assert ".docx" in payload["supported"]
    assert ".pdf" in payload["supported"]
